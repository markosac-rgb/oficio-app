"""Convierte docs/legal/*.md en .docx sencillos, sin inventar datos."""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
LEGAL = ROOT / "docs" / "legal"
PAIRS = (
    ("terminos.md", "terminos.docx"),
    ("aviso-privacidad.md", "aviso-privacidad.docx"),
)


def add_inline(paragraph, text: str) -> None:
    parts = text.split("**")
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1
        if part.startswith("[") and part.endswith("]"):
            run.bold = True
            run.font.highlight_color = 7


def convert(source: Path, target: Path) -> None:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)
    for raw in source.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line == "Borrador — debe revisarlo un abogado antes de publicarse":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.bold = True
            run.font.color.rgb = RGBColor(0x5C, 0x43, 0x10)
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].replace("**", ""), level=3)
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].replace("**", ""), level=2)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].replace("**", ""), level=1)
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, line)
    document.save(target)


def main() -> None:
    for source_name, target_name in PAIRS:
        convert(LEGAL / source_name, LEGAL / target_name)
        print(target_name)


if __name__ == "__main__":
    main()
